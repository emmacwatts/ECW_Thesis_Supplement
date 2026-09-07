#!/usr/bin/env python3
"""Create gene-expression plots from paired mean/error summary tables."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import numpy as np
import pandas as pd
from openpyxl import load_workbook

TIME_RE = re.compile(r"^(?P<series>.+)_(?P<time>\d+)$")


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Plot IDs selected by a mapping CSV from two expression summaries."
    )
    parser.add_argument("--main-summary", required=True, type=Path)
    parser.add_argument("--hamel-summary", required=True, type=Path)
    parser.add_argument("--gene-groups", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument(
        "--main-deseq-summary", type=Path,
        help="Optional DESeq2-normalised mean/error table for the time-course study",
    )
    parser.add_argument(
        "--hamel-deseq-summary", type=Path,
        help="Optional DESeq2-normalised mean/error table for the Hamel study",
    )
    parser.add_argument("--mock-series", default="wt_mock")
    parser.add_argument("--treatment-series", default="wt_ag")
    parser.add_argument("--dpi-label", default="Days post-inoculation")
    return parser.parse_args()


def read_table(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(f"Input file not found: {path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return pd.read_excel(path)
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    if path.suffix.lower() in {".tsv", ".txt"}:
        return pd.read_csv(path, sep="\t")
    raise ValueError(f"Unsupported table type for {path}; use CSV, TSV, or XLSX")


def load_summary(
    path: Path, requested_ids: set[str] | None = None
) -> tuple[pd.DataFrame, list[str]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"} and requested_ids is not None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(value) if value is not None else "" for value in next(rows)]
        except StopIteration as error:
            raise ValueError(f"{path} is empty") from error
        id_positions = [header.index(name) for name in ("gene_id", "target_id") if name in header]
        if not id_positions:
            raise ValueError(f"{path} needs a gene_id or target_id column")
        id_position = id_positions[0]
        selected = [row for row in rows if str(row[id_position]).strip() in requested_ids]
        workbook.close()
        frame = pd.DataFrame(selected, columns=header)
    else:
        frame = read_table(path)
    id_columns = [name for name in ("gene_id", "target_id") if name in frame]
    if not id_columns:
        raise ValueError(f"{path} needs a gene_id or target_id column")
    identifier = id_columns[0]
    frame[identifier] = frame[identifier].astype(str).str.strip()
    if frame[identifier].eq("").any() or frame[identifier].duplicated().any():
        raise ValueError(f"{path} contains blank or duplicate identifiers")
    frame = frame.set_index(identifier)
    conditions = [str(c)[:-5] for c in frame if str(c).endswith("_mean")]
    if not conditions:
        raise ValueError(f"{path} has no NAME_mean columns")
    missing = [f"{c}_error" for c in conditions if f"{c}_error" not in frame]
    if missing:
        raise ValueError(f"{path} is missing: {', '.join(missing)}")
    numeric = [column for c in conditions for column in (f"{c}_mean", f"{c}_error")]
    frame[numeric] = frame[numeric].apply(pd.to_numeric, errors="raise")
    if (frame[[f"{c}_error" for c in conditions]] < 0).any().any():
        raise ValueError(f"{path} contains negative error values")
    return frame, conditions


def safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._") or "target"


def load_groups(path: Path) -> pd.DataFrame:
    frame = read_table(path)
    missing = {"target", "gene_id"}.difference(frame)
    if missing:
        raise ValueError(f"{path} is missing: {', '.join(sorted(missing))}")
    frame = frame.copy()
    for column in ("target", "gene_id"):
        frame[column] = frame[column].astype(str).str.strip()
    if frame[["target", "gene_id"]].eq("").any().any():
        raise ValueError(f"{path} contains blank target or gene_id values")
    if frame.duplicated(["target", "gene_id"]).any():
        raise ValueError(f"{path} contains duplicate target/gene_id rows")
    if "image_file" not in frame:
        order = {target: i for i, target in enumerate(dict.fromkeys(frame.target), 1)}
        frame["image_file"] = [f"{order[t]:02d}_{safe_filename(t)}.png" for t in frame.target]
    else:
        frame["image_file"] = frame["image_file"].astype(str).str.strip()
    return frame


def series(conditions: list[str], prefix: str) -> list[tuple[int, str]]:
    found = []
    for condition in conditions:
        match = TIME_RE.match(condition)
        if match and match.group("series") == prefix:
            found.append((int(match.group("time")), condition))
    if not found:
        raise ValueError(f"No main-summary conditions match {prefix!r}")
    return sorted(found)


def plot_one(
    target: str,
    genes: list[str],
    main: pd.DataFrame,
    main_conditions: list[str],
    hamel: pd.DataFrame,
    hamel_conditions: list[str],
    mock_prefix: str,
    treatment_prefix: str,
    dpi_label: str,
) -> plt.Figure:
    plt.style.use("seaborn-v0_8-darkgrid")
    fig, axes = plt.subplots(1, 2, figsize=(14.5, 6.2))
    fig.subplots_adjust(top=0.73, left=0.06, right=0.98, bottom=0.15, wspace=0.36)

    preferred = ["Non-infiltrated", "Mock", "EV", "p19", "p19HA"]
    bar_conditions = [c for c in preferred if c in hamel_conditions]
    bar_conditions.extend(c for c in hamel_conditions if c not in bar_conditions)
    bar_colours = ["#aeb4ba", "#88a7c4", "#82b9a4", "#c59b87", "#9b8db6"]
    positions = np.arange(len(genes), dtype=float)
    width = 0.8 / len(bar_conditions)
    for i, condition in enumerate(bar_conditions):
        values = hamel.loc[genes, f"{condition}_mean"].to_numpy(float)
        errors = hamel.loc[genes, f"{condition}_error"].to_numpy(float)
        offset = (i - (len(bar_conditions) - 1) / 2) * width
        axes[0].bar(
            positions + offset, values, width, yerr=errors, capsize=3,
            label=condition, color=bar_colours[i % len(bar_colours)],
            edgecolor="#555555", linewidth=0.7, alpha=0.85,
        )
    axes[0].set_xticks(positions, genes)
    axes[0].set_xlabel("Gene ID")
    axes[0].set_ylabel("Mean TPM")
    axes[0].legend(
        loc="lower center", bbox_to_anchor=(0.5, 1.03), ncol=len(bar_conditions),
        frameon=False, fontsize=9,
    )

    palette = [
        ("#a9c2dc", "#5f8fbe", "o"),
        ("#9bcfbe", "#58a98d", "D"),
        ("#e0b9a5", "#ba7655", "s"),
        ("#c5b5d8", "#846ca3", "^"),
    ]
    for gene_index, gene in enumerate(genes):
        mock_colour, treatment_colour, marker = palette[gene_index % len(palette)]
        for prefix, label, colour in (
            (mock_prefix, "Mock", mock_colour),
            (treatment_prefix, "Agroinfiltrated", treatment_colour),
        ):
            selected = series(main_conditions, prefix)
            times = np.array([time for time, _ in selected])
            means = np.array([main.at[gene, f"{condition}_mean"] for _, condition in selected], float)
            sems = np.array([main.at[gene, f"{condition}_error"] for _, condition in selected], float)
            axes[1].plot(
                times, means, marker=marker, linewidth=2, color=colour,
                label=f"{gene} – {label}",
            )
            axes[1].fill_between(times, means - sems, means + sems, color=colour, alpha=0.18)
    axes[1].set_xlabel(dpi_label)
    axes[1].set_ylabel("Mean TPM")
    axes[1].legend(
        loc="lower center", bbox_to_anchor=(0.5, 1.03), ncol=2,
        frameon=False, fontsize=8,
    )
    for axis in axes:
        axis.set_ylim(bottom=0)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    fig.suptitle(target, x=0.06, y=0.98, ha="left", fontsize=20, fontstyle="italic")
    return fig


def plot_deseq2_target(
    target: str,
    genes: list[str],
    main: pd.DataFrame,
    main_conditions: list[str],
    hamel: pd.DataFrame,
    hamel_conditions: list[str],
    mock_prefix: str,
    treatment_prefix: str,
    dpi_label: str,
) -> plt.Figure:
    """Plot each gene separately so all comparisons remain within one gene."""
    plt.style.use("seaborn-v0_8-darkgrid")
    fig, axes = plt.subplots(
        len(genes), 2, figsize=(14.5, 4.6 * len(genes)), squeeze=False,
    )
    fig.subplots_adjust(
        top=0.88 if len(genes) == 1 else 0.93,
        left=0.07, right=0.98, bottom=0.09, hspace=0.48, wspace=0.30,
    )
    preferred = ["Non-infiltrated", "Mock", "EV", "p19", "p19HA"]
    bar_conditions = [c for c in preferred if c in hamel_conditions]
    bar_conditions.extend(c for c in hamel_conditions if c not in bar_conditions)
    bar_colours = ["#aeb4ba", "#88a7c4", "#82b9a4", "#c59b87", "#9b8db6"]

    for row_index, gene in enumerate(genes):
        bar_axis, line_axis = axes[row_index]
        positions = np.arange(len(bar_conditions))
        means = hamel.loc[gene, [f"{c}_mean" for c in bar_conditions]].to_numpy(float)
        errors = hamel.loc[gene, [f"{c}_error" for c in bar_conditions]].to_numpy(float)
        bar_axis.bar(
            positions, means, yerr=errors, capsize=3,
            color=bar_colours[:len(bar_conditions)], edgecolor="#555555",
            linewidth=0.7, alpha=0.85,
        )
        bar_axis.set_xticks(positions, bar_conditions, rotation=25, ha="right")
        bar_axis.set_title(f"{gene} – Hamel et al.")
        bar_axis.set_ylabel("Mean DESeq2-normalised count")

        for prefix, label, colour in (
            (mock_prefix, "Mock", "#a9c2dc"),
            (treatment_prefix, "Agroinfiltrated", "#5f8fbe"),
        ):
            selected = series(main_conditions, prefix)
            times = np.array([time for time, _ in selected])
            means = np.array([main.at[gene, f"{condition}_mean"] for _, condition in selected], float)
            errors = np.array([main.at[gene, f"{condition}_error"] for _, condition in selected], float)
            line_axis.plot(times, means, marker="o", linewidth=2, color=colour, label=label)
            line_axis.fill_between(times, means - errors, means + errors, color=colour, alpha=0.18)
        line_axis.set_title(f"{gene} – Grosse-Holz et al.")
        line_axis.set_xlabel(dpi_label)
        line_axis.set_ylabel("Mean DESeq2-normalised count")
        line_axis.legend(frameon=False)

        for axis in (bar_axis, line_axis):
            axis.set_ylim(bottom=0)
            axis.spines["top"].set_visible(False)
            axis.spines["right"].set_visible(False)

    fig.suptitle(target, x=0.07, y=0.99, ha="left", fontsize=20, fontstyle="italic")
    return fig


def values_long(
    groups: pd.DataFrame,
    study: str,
    summary: pd.DataFrame,
    conditions: list[str],
) -> list[dict[str, object]]:
    return [
        {
            "target": row.target,
            "gene_id": row.gene_id,
            "study": study,
            "condition": condition,
            "mean": summary.at[row.gene_id, f"{condition}_mean"],
            "error": summary.at[row.gene_id, f"{condition}_error"],
        }
        for row in groups.itertuples(index=False)
        for condition in conditions
    ]


def main() -> None:
    args = arguments()
    if (args.main_deseq_summary is None) != (args.hamel_deseq_summary is None):
        raise ValueError(
            "Supply both --main-deseq-summary and --hamel-deseq-summary, or neither"
        )
    groups = load_groups(args.gene_groups)
    requested_ids = set(groups.gene_id)
    main_table, main_conditions = load_summary(args.main_summary, requested_ids)
    hamel_table, hamel_conditions = load_summary(args.hamel_summary, requested_ids)
    for label, table in (("main", main_table), ("Hamel", hamel_table)):
        missing = sorted(set(groups.gene_id).difference(table.index))
        if missing:
            raise ValueError(f"IDs missing from {label} summary: {', '.join(missing)}")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    with PdfPages(args.output_dir / "gene_expression_plots.pdf") as pdf:
        count = 0
        for target, target_rows in groups.groupby("target", sort=False):
            genes = target_rows.gene_id.tolist()
            filename = target_rows.image_file.iloc[0]
            filename = filename if filename.lower().endswith(".png") else f"{filename}.png"
            figure = plot_one(
                target, genes, main_table, main_conditions, hamel_table,
                hamel_conditions, args.mock_series, args.treatment_series, args.dpi_label,
            )
            figure.savefig(args.output_dir / filename, dpi=300, bbox_inches="tight")
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)
            count += 1

    rows = values_long(groups, "main", main_table, main_conditions)
    rows.extend(values_long(groups, "hamel", hamel_table, hamel_conditions))
    pd.DataFrame(rows).to_csv(args.output_dir / "gene_expression_values.csv", index=False)
    print(f"Created {count} target plots in {args.output_dir}")

    if args.main_deseq_summary is not None:
        main_deseq, main_deseq_conditions = load_summary(
            args.main_deseq_summary, requested_ids
        )
        hamel_deseq, hamel_deseq_conditions = load_summary(
            args.hamel_deseq_summary, requested_ids
        )
        for label, table in (("main DESeq2", main_deseq), ("Hamel DESeq2", hamel_deseq)):
            missing = sorted(requested_ids.difference(table.index))
            if missing:
                raise ValueError(f"IDs missing from {label} summary: {', '.join(missing)}")

        deseq_dir = args.output_dir / "deseq2"
        deseq_dir.mkdir(parents=True, exist_ok=True)
        deseq_count = 0
        with PdfPages(deseq_dir / "gene_expression_plots_DESeq2.pdf") as pdf:
            for target, target_rows in groups.groupby("target", sort=False):
                genes = target_rows.gene_id.tolist()
                base = Path(target_rows.image_file.iloc[0]).stem
                figure = plot_deseq2_target(
                    target, genes, main_deseq, main_deseq_conditions,
                    hamel_deseq, hamel_deseq_conditions,
                    args.mock_series, args.treatment_series, args.dpi_label,
                )
                figure.savefig(
                    deseq_dir / f"{base}_DESeq2.png", dpi=300, bbox_inches="tight"
                )
                pdf.savefig(figure, bbox_inches="tight")
                plt.close(figure)
                deseq_count += 1

        deseq_rows = values_long(
            groups, "main", main_deseq, main_deseq_conditions
        )
        deseq_rows.extend(values_long(
            groups, "hamel", hamel_deseq, hamel_deseq_conditions
        ))
        pd.DataFrame(deseq_rows).to_csv(
            deseq_dir / "gene_expression_values_DESeq2.csv", index=False
        )
        print(f"Created {deseq_count} DESeq2 target figures in {deseq_dir}")


if __name__ == "__main__":
    main()
