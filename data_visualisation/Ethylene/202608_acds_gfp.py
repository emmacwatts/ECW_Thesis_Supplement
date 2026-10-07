"""Plot the acdS GFP workbook."""

from __future__ import annotations

import math
import os
from pathlib import Path

os.environ.setdefault("MPLBACKEND", "Agg")
os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parents[1] / ".matplotlib-cache"))

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats
from statsmodels.stats.multitest import multipletests


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "originalData" / "acdS dataset.xlsx"
FIGURES = HERE / "202608_acds_gfp_figures"
LOGS = HERE / "202608_acds_gfp_logs"

FONT_STACK = ["Helvetica Neue", "Helvetica", "Arial", "DejaVu Sans"]
TEXT = "#4A4A4A"
PANEL_BG = "#E9EFF6"
WT_COLOR = "#4E79A7"
ACDS_COLOR = "#B07AA1"  # muted purple, distinct from 40-1 orange and ein2 green
WT_POINT_COLOR = "#2F5F8F"
ACDS_POINT_COLOR = "#85577E"
BAR_ALPHA = 0.76
POINT_ALPHA = 0.90

plt.rcParams.update({
    "font.family": "sans-serif", "font.sans-serif": FONT_STACK,
    "font.size": 14, "axes.labelsize": 14, "xtick.labelsize": 13,
    "ytick.labelsize": 13, "legend.fontsize": 13,
    "axes.labelcolor": TEXT, "axes.edgecolor": TEXT, "axes.linewidth": 1.0,
    "xtick.color": TEXT, "ytick.color": TEXT, "text.color": TEXT,
    "svg.fonttype": "none", "pdf.fonttype": 42,
    "figure.dpi": 120, "savefig.dpi": 300,
})


def workbook_values():
    return load_workbook(SOURCE, data_only=False).active, load_workbook(SOURCE, data_only=True).active


def independently_process(values):
    """Use only raw cells C:E in the first table; do not read processed cells F:G."""
    records = []
    for row in range(4, 14):
        genotype = "EV" if row <= 9 else "acdS"
        experiment = "1" if row <= 6 or row >= 10 else "2"
        within_experiment = row - 3 if row <= 6 else row - 6 if row <= 9 else row - 9
        sample = f"{genotype}_rep{experiment}_sample{within_experiment}"
        background = float(values[f"E{row}"].value)
        for p19, column in (("−", "D"), ("+", "C")):
            raw = float(values[f"{column}{row}"].value)
            records.append({
                "source_row": row, "sample": sample, "genotype": genotype,
                "experiment": experiment, "p19": p19, "raw_gfp": raw,
                "background": background, "background_corrected_gfp": raw - background,
            })
    data = pd.DataFrame(records)
    data["genotype"] = pd.Categorical(data["genotype"], ["EV", "acdS"], ordered=True)
    data["p19"] = pd.Categorical(data["p19"], ["−", "+"], ordered=True)
    return data


def audit_second_table(formulas, values, data):
    """Compare independent results to the workbook processing only after calculation."""
    checks = []
    processed_map = {"+": "F", "−": "G"}
    second_map = {
        ("EV", "+"): ("C", range(18, 24)), ("EV", "−"): ("D", range(18, 24)),
        ("acdS", "+"): ("E", range(18, 22)), ("acdS", "−"): ("F", range(18, 22)),
    }
    for record in data.to_dict("records"):
        p19 = str(record["p19"])
        expected = record["background_corrected_gfp"]
        displayed = float(values[f"{processed_map[p19]}{record['source_row']}"].value)
        checks.append({"check": "first-table subtraction", "group": f"{record['genotype']} {p19}p19",
                       "cell": f"{processed_map[p19]}{record['source_row']}", "independent": expected,
                       "workbook": displayed, "difference": displayed - expected,
                       "matches": bool(np.isclose(displayed, expected))})

    for (genotype, p19), (column, target_rows) in second_map.items():
        independent = data.loc[(data["genotype"] == genotype) & (data["p19"] == p19),
                               "background_corrected_gfp"].to_numpy(float)
        for expected, row in zip(independent, target_rows):
            displayed = float(values[f"{column}{row}"].value)
            checks.append({"check": "second-table transfer", "group": f"{genotype} {p19}p19",
                           "cell": f"{column}{row}", "independent": expected,
                           "workbook": displayed, "difference": displayed - expected,
                           "matches": bool(np.isclose(displayed, expected))})
        summary = {25: independent.mean(), 26: independent.std(ddof=1),
                   27: independent.std(ddof=1) / math.sqrt(len(independent))}
        for row, expected in summary.items():
            displayed = float(values[f"{column}{row}"].value)
            checks.append({"check": {25: "mean", 26: "sample SD", 27: "SEM"}[row],
                           "group": f"{genotype} {p19}p19", "cell": f"{column}{row}",
                           "formula": formulas[f"{column}{row}"].value, "independent": expected,
                           "workbook": displayed, "difference": displayed - expected,
                           "matches": bool(np.isclose(displayed, expected, rtol=1e-7, atol=1e-7))})
    return pd.DataFrame(checks)


def analyse(data):
    genotype_rows = []
    for p19 in ("−", "+"):
        wt = data.loc[(data["genotype"] == "EV") & (data["p19"] == p19), "background_corrected_gfp"]
        acds = data.loc[(data["genotype"] == "acdS") & (data["p19"] == p19), "background_corrected_gfp"]
        result = stats.ttest_ind(wt, acds, equal_var=False)
        genotype_rows.append({"p19": p19, "comparison": "EV vs acdS",
                              "test": "two-sided Welch t-test", "wt_n": len(wt), "acds_n": len(acds),
                              "wt_mean": wt.mean(), "acds_mean": acds.mean(), "difference_wt_minus_acds": wt.mean()-acds.mean(),
                              "t": result.statistic, "df": result.df, "p_raw": result.pvalue})
    adjusted = multipletests([row["p_raw"] for row in genotype_rows], method="holm")[1]
    for row, p_value in zip(genotype_rows, adjusted):
        row["p_holm"] = p_value
        row["significant_p_holm_lt_0.05"] = bool(p_value < 0.05)

    paired_rows, changes = [], {}
    for genotype in ("EV", "acdS"):
        wide = data.loc[data["genotype"] == genotype].pivot(index="sample", columns="p19", values="background_corrected_gfp")
        result = stats.ttest_rel(wide["+"], wide["−"])
        changes[genotype] = (wide["+"] - wide["−"]).to_numpy(float)
        paired_rows.append({"genotype": genotype, "comparison": "+p19 vs −p19",
                            "test": "two-sided paired t-test", "n_pairs": len(wide),
                            "mean_paired_change": changes[genotype].mean(), "t": result.statistic,
                            "df": result.df, "p_raw": result.pvalue})
    paired_adjusted = multipletests([row["p_raw"] for row in paired_rows], method="holm")[1]
    for row, p_value in zip(paired_rows, paired_adjusted):
        row["p_holm"] = p_value
        row["significant_p_holm_lt_0.05"] = bool(p_value < 0.05)

    interaction = stats.ttest_ind(changes["EV"], changes["acdS"], equal_var=False)
    interaction_row = pd.DataFrame([{
        "comparison": "EV vs acdS in paired (+p19 minus −p19) change",
        "test": "two-sided Welch difference-of-differences t-test",
        "wt_n": len(changes["EV"]), "acds_n": len(changes["acdS"]),
        "wt_mean_change": changes["EV"].mean(), "acds_mean_change": changes["acdS"].mean(),
        "t": interaction.statistic, "df": interaction.df, "p_raw": interaction.pvalue,
    }])
    return pd.DataFrame(genotype_rows), pd.DataFrame(paired_rows), interaction_row


def p_label(p_value):
    if p_value < 0.001: return "***"
    if p_value < 0.01: return "**"
    if p_value < 0.05: return "*"
    return "ns"


def significance(ax, left, right, y, text):
    height = 0.55
    ax.plot([left, left, right, right], [y, y + height, y + height, y], color="black", linewidth=1.0)
    ax.text((left + right) / 2, y + height + 0.17, text, ha="center", va="bottom", fontsize=12, color=TEXT)


def make_plot(data, contrasts):
    fig, ax = plt.subplots(figsize=(8.2, 4.4), facecolor="white")
    ax.set_facecolor(PANEL_BG); ax.set_axisbelow(True)
    ax.grid(axis="y", color="white", linewidth=0.9, alpha=0.9)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(width=1.0, length=4.5, direction="out")
    centers = np.arange(2, dtype=float)
    offsets = {"EV": -0.18, "acdS": 0.18}
    colors = {"EV": WT_COLOR, "acdS": ACDS_COLOR}
    point_colors = {"EV": WT_POINT_COLOR, "acdS": ACDS_POINT_COLOR}
    for p19, center in zip(("−", "+"), centers):
        for genotype in ("EV", "acdS"):
            values = data.loc[(data["p19"] == p19) & (data["genotype"] == genotype),
                              "background_corrected_gfp"].to_numpy(float) / 1000
            xpos = center + offsets[genotype]
            mean, sem = values.mean(), values.std(ddof=1) / math.sqrt(len(values))
            ax.bar(xpos, mean, width=0.32, color=colors[genotype], alpha=BAR_ALPHA,
                   edgecolor="black", linewidth=0.9, zorder=2)
            ax.errorbar(xpos, mean, yerr=sem, fmt="none", ecolor="black", elinewidth=1,
                        capsize=3, capthick=1, zorder=4)
            jitter = np.linspace(-0.045, 0.045, len(values))
            ax.scatter(xpos + jitter, values, s=24, color=point_colors[genotype],
                       edgecolor="#000000", linewidth=0.45,
                       alpha=POINT_ALPHA, zorder=3)
    for center, p19, y in zip(centers, ("−", "+"), (8.8, 20.5)):
        p_value = contrasts.loc[contrasts["p19"] == p19, "p_holm"].iloc[0]
        significance(ax, center - 0.18, center + 0.18, y, p_label(p_value))
    ax.set_xticks(centers, ["−", "+"]); ax.set_xlabel("p19")
    ax.set_ylabel("Normalised GFP intensity (AU × 10³)"); ax.set_ylim(0, 23)
    handles = [plt.Rectangle((0, 0), 1, 1, facecolor=colors[label], edgecolor="black",
                             linewidth=0.8, alpha=BAR_ALPHA, label=label) for label in ("EV", "acdS")]
    legend = ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, -0.18),
                       ncol=2, frameon=False, handlelength=1.2, columnspacing=1.3)
    legend.get_texts()[1].set_fontstyle("italic")
    fig.subplots_adjust(left=0.18, right=0.98, top=0.95, bottom=0.24)
    FIGURES.mkdir(parents=True, exist_ok=True)
    for extension in ("png", "pdf", "svg"):
        fig.savefig(FIGURES / f"202608_acds_gfp_summary.{extension}", dpi=300,
                    bbox_inches="tight", facecolor="white")
    plt.close(fig)


def write_outputs(data, audit, contrasts, paired, interaction):
    LOGS.mkdir(parents=True, exist_ok=True)
    data.to_csv(LOGS / "acds_background_corrected_long.csv", index=False)
    audit.to_csv(LOGS / "acds_workbook_processing_audit.csv", index=False)
    contrasts.to_csv(LOGS / "acds_genotype_contrasts.csv", index=False)
    paired.to_csv(LOGS / "acds_p19_paired_tests.csv", index=False)
    interaction.to_csv(LOGS / "acds_genotype_by_p19_interaction.csv", index=False)
    minus = contrasts.loc[contrasts["p19"] == "−"].iloc[0]
    plus = contrasts.loc[contrasts["p19"] == "+"].iloc[0]
    wt_paired = paired.loc[paired["genotype"] == "EV"].iloc[0]
    acds_paired = paired.loc[paired["genotype"] == "acdS"].iloc[0]
    note = f"""acdS GFP figure: processing, statistics, audit, and figure-legend notes

Descriptive results (mean +/- SEM, AU)
EV -p19: {minus['wt_mean']:.3f} +/- {data.loc[(data['genotype']=='EV') & (data['p19']=='−'),'background_corrected_gfp'].sem():.3f} (n={int(minus['wt_n'])}).
acdS -p19: {minus['acds_mean']:.3f} +/- {data.loc[(data['genotype']=='acdS') & (data['p19']=='−'),'background_corrected_gfp'].sem():.3f} (n={int(minus['acds_n'])}).
EV +p19: {plus['wt_mean']:.3f} +/- {data.loc[(data['genotype']=='EV') & (data['p19']=='+'),'background_corrected_gfp'].sem():.3f} (n={int(plus['wt_n'])}).
acdS +p19: {plus['acds_mean']:.3f} +/- {data.loc[(data['genotype']=='acdS') & (data['p19']=='+'),'background_corrected_gfp'].sem():.3f} (n={int(plus['acds_n'])}).

"""
   
def main():
    formulas, values = workbook_values()
    data = independently_process(values)
    audit = audit_second_table(formulas, values, data)
    if not audit["matches"].all():
        raise ValueError("Workbook cross-check failed; inspect acds_workbook_processing_audit.csv")
    contrasts, paired, interaction = analyse(data)
    make_plot(data, contrasts)
    write_outputs(data, audit, contrasts, paired, interaction)


if __name__ == "__main__":
    main()
