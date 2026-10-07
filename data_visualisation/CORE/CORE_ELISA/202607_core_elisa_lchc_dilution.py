"""CORE LCHC ELISA dilution curves across Cas9 WT and CORE mutant plants."""

from pathlib import Path
import os
import sys

MPLCONFIGDIR = Path.cwd() / ".matplotlib-cache"
MPLCONFIGDIR.mkdir(exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(MPLCONFIGDIR))
os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats


HERE = Path(__file__).resolve().parent
STYLE_DIR = HERE.parents[1] / "PRp27" / "PRp27_timecourse_202606"
if str(STYLE_DIR) not in sys.path:
    sys.path.insert(0, str(STYLE_DIR))
from figure_styles import ANALYSIS_RCPARAMS, FONT_STACK, TEXT_COLOR, apply_axis_style  # noqa: E402

SOURCE = Path(
    "/Users/user/Library/CloudStorage/OneDrive-Nexus365/MolecularFarming/"
    "6. PosHits/CORE/GFP-COVA-VirB/CombinedTest/COVA_ELISA_6wk_CORECombined.xlsx"
)
SHEET = "Result sheet (3)"
SCRIPT_STEM = Path(__file__).stem
SAVE_DIR = HERE / f"{SCRIPT_STEM}_figures"
LOG_DIR = HERE / f"{SCRIPT_STEM}_logs"
for folder in [SAVE_DIR, LOG_DIR]:
    folder.mkdir(parents=True, exist_ok=True)

GENOTYPE_ORDER = ["Cas9 WT", "core-714-1", "core-704-6", "core-2-8", "core-1-2"]
COLORS = {
    "Cas9 WT": "#79C6A3",
    "core-714-1": "#38A6A5",
    "core-704-6": "#287FB8",
    "core-2-8": "#536FA8",
    "core-1-2": "#31539A",
}
MARKERS = {"Cas9 WT": "o", "core-714-1": "s", "core-704-6": "^",
           "core-2-8": "D", "core-1-2": "P"}
SIGNIFICANCE_LETTERS = {
    "core-714-1": "a", "core-704-6": "b", "core-2-8": "c", "core-1-2": "d",
}

# Columns are two biological replicates per genotype.
COLUMN_MAP = {
    1: ("Cas9 WT", 1), 2: ("Cas9 WT", 2),
    3: ("core-1-2", 1), 4: ("core-1-2", 2),
    5: ("core-2-8", 1), 6: ("core-2-8", 2),
    7: ("core-704-6", 1), 8: ("core-704-6", 2),
    9: ("core-705-1", 1), 10: ("core-705-1", 2),
    11: ("core-714-1", 1), 12: ("core-714-1", 2),
}
# Successive row pairs are technical replicates at each dilution.
DILUTION_MAP = {
    "1:3": ("A", "B"), "1:30": ("C", "D"),
    "1:300": ("E", "F"), "1:3000": ("G", "H"),
}
DILUTION_ORDER = ["1:3000", "1:300", "1:30", "1:3"]
DILUTION_FRACTION = {label: 1 / float(label.split(":")[1]) for label in DILUTION_ORDER}
FLAGGED_WELLS = {"D1", "D10"}
EXCLUDED_GENOTYPES = {"core-705-1"}

METHODS = (
    "LCHC ELISA absorbance values were read from COVA_ELISA_6wk_CORECombined.xlsx. "
    "Columns 1-2, 3-4, 5-6, 7-8, 9-10 and 11-12 represent two biological replicates "
    "of Cas9 WT, core-1-2, core-2-8, core-704-6, core-705-1 and core-714-1, "
    "respectively. Row pairs A-B, C-D, E-F and G-H are technical replicates at 1:3, "
    "1:30, 1:300 and 1:3000. Technical replicates were averaged before biological "
    "summaries. D1 and D10 were excluded as flagged in the source workbook; the C1 "
    "measurement therefore provides the 1:30 value for WT biological replicate 1. "
    "CORE 705-1 was excluded entirely. Curves show mean +/- SD across biological "
    "replicates as translucent ribbons, with individual biological-replicate means "
    "overlaid. At each dilution, every CORE line was compared with Cas9 WT using a "
    "two-sided Dunnett multiple-comparison test, which simultaneously adjusts the four "
    "treatment-versus-control comparisons. Significant adjusted comparisons (p < 0.05) "
    "are marked above the dilution cluster "
    "using fixed colour-matched genotype letters: a, core-714-1; b, core-704-6; c, "
    "core-2-8; d, core-1-2. No background subtraction, transformation, or additional "
    "outlier removal was applied."
)

plt.rcParams.update({
    **ANALYSIS_RCPARAMS,
    "font.family": "sans-serif", "font.sans-serif": FONT_STACK,
    "font.size": 17, "axes.titlesize": 19, "axes.labelsize": 18,
    "xtick.labelsize": 17, "ytick.labelsize": 17, "legend.fontsize": 16,
    "figure.dpi": 120, "savefig.dpi": 300,
})


def read_wells():
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    sheet = workbook[SHEET]
    values = {str(sheet.cell(row, 1).value): sheet.cell(row, 2).value for row in range(42, 138)}
    rows = []
    for dilution, technical_rows in DILUTION_MAP.items():
        for column, (genotype, biological_rep) in COLUMN_MAP.items():
            for technical_rep, row_name in enumerate(technical_rows, start=1):
                well = f"{row_name}{column}"
                rows.append({
                    "well": well, "genotype": genotype, "biological_replicate": biological_rep,
                    "technical_replicate": technical_rep, "dilution": dilution,
                    "dilution_fraction": DILUTION_FRACTION[dilution],
                    "a450": pd.to_numeric(values.get(well), errors="coerce"),
                    "flagged_source_exclusion": well in FLAGGED_WELLS,
                    "excluded_genotype": genotype in EXCLUDED_GENOTYPES,
                })
    return pd.DataFrame(rows)


def process(raw):
    included = raw[~raw["flagged_source_exclusion"] & ~raw["excluded_genotype"]].copy()
    biological = (
        included.groupby(["genotype", "biological_replicate", "dilution", "dilution_fraction"])
        .agg(a450=("a450", "mean"), n_technical=("a450", "count"),
             technical_sd=("a450", "std"), wells=("well", lambda x: ", ".join(x)))
        .reset_index()
    )
    summary = (
        biological.groupby(["genotype", "dilution", "dilution_fraction"])
        .agg(n_biological=("biological_replicate", "nunique"), mean_a450=("a450", "mean"),
             sd_a450=("a450", "std"), sem_a450=("a450", "sem"))
        .reset_index()
    )
    return included, biological, summary


def genotype_vs_wt_tests(biological):
    rows = []
    for dilution in DILUTION_ORDER:
        at_dilution = biological[biological["dilution"].eq(dilution)]
        wt = at_dilution.loc[at_dilution["genotype"].eq("Cas9 WT"), "a450"].to_numpy()
        core_values = [
            at_dilution.loc[at_dilution["genotype"].eq(genotype), "a450"].to_numpy()
            for genotype in GENOTYPE_ORDER[1:]
        ]
        result = stats.dunnett(
            *core_values, control=wt, alternative="two-sided", random_state=202607
        )
        for genotype, core, statistic, p_adjusted in zip(
            GENOTYPE_ORDER[1:], core_values, result.statistic, result.pvalue
        ):
            rows.append({
                "dilution": dilution, "genotype": genotype,
                "letter": SIGNIFICANCE_LETTERS[genotype],
                "n_wt": len(wt), "n_core": len(core),
                "mean_core_minus_wt": np.mean(core) - np.mean(wt),
                "dunnett_statistic": statistic,
                "p_dunnett_adjusted": p_adjusted,
                "test": "two-sided Dunnett test versus Cas9 WT",
            })
    table = pd.DataFrame(rows)
    table["significant_0.05"] = table["p_dunnett_adjusted"] < 0.05
    return table


def style_axis(ax):
    ax.set_facecolor("#E9EFF6")
    ax.grid(True, color="white", linewidth=0.9, alpha=0.9)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    apply_axis_style(ax)
    for spine in ax.spines.values():
        spine.set_color(TEXT_COLOR)


def plot(biological, summary, tests):
    fig, ax = plt.subplots(figsize=(10.4, 7.0))
    x_positions = np.arange(len(DILUTION_ORDER), dtype=float)
    rng = np.random.default_rng(202607)
    for genotype in GENOTYPE_ORDER:
        stats_df = summary[summary["genotype"].eq(genotype)].set_index("dilution").loc[DILUTION_ORDER]
        raw = biological[biological["genotype"].eq(genotype)]
        means = stats_df["mean_a450"].to_numpy()
        errors = stats_df["sd_a450"].to_numpy()
        color = COLORS[genotype]
        for index, dilution in enumerate(DILUTION_ORDER):
            values = raw.loc[raw["dilution"].eq(dilution), "a450"].to_numpy()
            jitter = rng.uniform(-0.055, 0.055, len(values))
            ax.scatter(index + jitter, values, s=31, color=color, alpha=0.55,
                       edgecolor="none", zorder=3)
        ax.fill_between(x_positions, np.maximum(0, means - errors), means + errors,
                        color=color, alpha=0.14, linewidth=0, zorder=2)
        ax.plot(x_positions, means, color=color, marker=MARKERS[genotype],
                markersize=7, markeredgecolor="white", markeredgewidth=0.7,
                linewidth=2.0, label=genotype, zorder=4)

    # Put all significant genotype letters together above each full dilution
    # cluster, rather than attaching annotations to closely spaced curves.
    raw_max = biological.groupby("dilution")["a450"].max()
    band_max = (summary.assign(upper=summary["mean_a450"] + summary["sd_a450"])
                .groupby("dilution")["upper"].max())
    global_top = max(raw_max.max(), band_max.max())
    letter_y = {}
    for index, dilution in enumerate(DILUTION_ORDER):
        y = max(raw_max[dilution], band_max[dilution]) + 0.055 * global_top
        letter_y[dilution] = y
        significant = tests[(tests["dilution"].eq(dilution)) & tests["significant_0.05"]]
        letters = significant["letter"].tolist()
        if letters:
            vertical_step = 0.045 * global_top
            for letter_index, (_, result) in enumerate(significant.iterrows()):
                stacked_y = y + (len(letters) - 1 - letter_index) * vertical_step
                ax.text(index, stacked_y, result["letter"], ha="center", va="bottom",
                        fontsize=17, fontweight="bold", color=COLORS[result["genotype"]])
        else:
            ax.text(index, y, "ns", ha="center", va="bottom", fontsize=14,
                    color=TEXT_COLOR)

    ax.set_xticks(x_positions, DILUTION_ORDER)
    ax.set_xlabel("Leaf extract dilution")
    ax.set_ylabel(r"Absorbance ($A_{450}$)")
    ax.set_ylim(0, max(letter_y.values()) + 0.09 * global_top)
    style_axis(ax)
    legend_labels = [
        "Cas9 WT",
        r"$\it{core}$-714-1 (a)", r"$\it{core}$-704-6 (b)",
        r"$\it{core}$-2-8 (c)", r"$\it{core}$-1-2 (d)",
    ]
    handles, _ = ax.get_legend_handles_labels()
    ax.legend(handles, legend_labels, frameon=False,
              loc="upper left", bbox_to_anchor=(1.01, 1.0))
    fig.subplots_adjust(left=0.13, right=0.76, top=0.96, bottom=0.15)
    for fmt in ["png", "pdf", "svg"]:
        fig.savefig(SAVE_DIR / f"{SCRIPT_STEM}_all_genotypes_dilution_curve.{fmt}",
                    dpi=300, bbox_inches="tight")
    plt.close(fig)


def main():
    raw = read_wells()
    included, biological, summary = process(raw)
    tests = genotype_vs_wt_tests(biological)
    raw.to_excel(LOG_DIR / "core_elisa_plate_map_all_wells.xlsx", index=False)
    included.to_excel(LOG_DIR / "core_elisa_included_technical_wells.xlsx", index=False)
    biological.to_excel(LOG_DIR / "core_elisa_biological_replicate_means.xlsx", index=False)
    summary.to_excel(LOG_DIR / "core_elisa_summary_mean_sd.xlsx", index=False)
    tests.to_excel(LOG_DIR / "core_elisa_genotype_vs_wt_tests.xlsx", index=False)
plot(biological, summary, tests)
    print(summary.to_string(index=False))
    print(f"Figures saved to: {SAVE_DIR}")
    print(tests.to_string(index=False))
    return raw, included, biological, summary, tests


if __name__ == "__main__":
    main()
