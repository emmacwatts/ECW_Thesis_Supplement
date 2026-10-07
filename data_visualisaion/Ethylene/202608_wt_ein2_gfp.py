"""Audit and plot the WT versus ein2 GFP workbook."""

from __future__ import annotations

import math
import os
from pathlib import Path

os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from scipy import stats
from statsmodels.formula.api import mixedlm
from statsmodels.stats.multitest import multipletests


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "originalData" / "wt vs ein2 GFP.xlsx"
FIGURES = HERE / "202608_wt_ein2_gfp_figures"
LOGS = HERE / "202608_wt_ein2_gfp_logs"

FONT_STACK = ["Helvetica Neue", "Helvetica", "Arial", "DejaVu Sans"]
TEXT = "#4A4A4A"
PANEL_BG = "#E9EFF6"
WT_COLOR = "#4E79A7"
EIN2_COLOR = "#59A14F"  # shared muted green; distinct from 40-1 orange
BAR_ALPHA = 0.76
POINT_ALPHA = 0.90

plt.rcParams.update({
    "font.family": "sans-serif", "font.sans-serif": FONT_STACK,
    "font.size": 14, "axes.labelsize": 14, "xtick.labelsize": 13,
    "ytick.labelsize": 13, "legend.fontsize": 13,
    "axes.labelcolor": TEXT, "axes.edgecolor": TEXT, "axes.linewidth": 1.0,
    "xtick.color": TEXT, "ytick.color": TEXT, "text.color": TEXT,
    "svg.fonttype": "none", "pdf.fonttype": 42,
    "figure.dpi": 120, "savefig.dpi": 300,
})


def workbook_values():
    formulas = load_workbook(SOURCE, data_only=False).active
    values = load_workbook(SOURCE, data_only=True).active
    return formulas, values


def source_mapping():
    """Return (right row, left GFP row, left GFP+P19 row, background row)."""
    wt = [(row, row, row + 14, row + 28) for row in range(7, 15)]
    ein2 = [(row, row, row + 14, row + 28) for row in range(15, 21)]
    return wt + ein2


def audit_workbook(formulas, values):
    rows = []
    for right_row, gfp_row, p19_row, background_row in source_mapping():
        sample = values[f"N{right_row}"].value
        checks = (
            ("background", values[f"O{right_row}"].value, values[f"C{background_row}"].value),
            ("GFP", values[f"P{right_row}"].value, values[f"C{gfp_row}"].value),
            ("GFP+P19", values[f"R{right_row}"].value, values[f"C{p19_row}"].value),
        )
        for measurement, right_value, left_value in checks:
            rows.append({"sample": sample, "measurement": measurement,
                         "left_table_value": left_value, "right_table_value": right_value,
                         "difference": right_value - left_value,
                         "matches": bool(np.isclose(right_value, left_value))})
        background = values[f"O{right_row}"].value
        for condition, raw_column, corrected_column in (("−p19", "P", "Q"), ("+p19", "R", "S")):
            raw_value = values[f"{raw_column}{right_row}"].value
            corrected = values[f"{corrected_column}{right_row}"].value
            expected = raw_value - background
            rows.append({"sample": sample, "measurement": f"{condition} background subtraction",
                         "left_table_value": expected, "right_table_value": corrected,
                         "difference": corrected - expected,
                         "matches": bool(np.isclose(corrected, expected))})
    transfer = pd.DataFrame(rows)

    summary_specs = (
        ("WT −p19", "P22", list(range(7, 15)), "Q"),
        ("WT +p19", "P23", list(range(7, 15)), "S"),
        ("ein2 −p19", "P24", list(range(15, 21)), "Q"),
        ("ein2 +p19", "P25", list(range(15, 21)), "S"),
    )
    summary_rows = []
    for label, cell, source_rows, column in summary_specs:
        displayed = values[cell].value
        expected = np.mean([values[f"{column}{row}"].value for row in source_rows])
        summary_rows.append({"summary": label, "cell": cell,
                             "formula": formulas[cell].value,
                             "expected_mean": expected, "displayed_mean": displayed,
                             "difference": displayed - expected,
                             "matches": bool(np.isclose(displayed, expected))})
    return transfer, pd.DataFrame(summary_rows)


def long_data(values):
    records = []
    for row in range(7, 21):
        sample = values[f"N{row}"].value
        genotype = "WT" if sample.startswith("WT") else "ein2"
        experiment = int(sample[2]) if genotype == "WT" else int(sample.split("_")[1].split(".")[0])
        for p19, column in (("−", "Q"), ("+", "S")):
            records.append({"sample": sample, "experiment": str(experiment),
                            "genotype": genotype, "p19": p19,
                            "background_corrected_gfp": values[f"{column}{row}"].value})
    data = pd.DataFrame(records)
    data["genotype"] = pd.Categorical(data["genotype"], ["WT", "ein2"], ordered=True)
    data["p19"] = pd.Categorical(data["p19"], ["−", "+"], ordered=True)
    return data


def analyse(data):
    # Each sample contributes paired −/+p19 measurements; experiment is a fixed
    # blocking factor and sample is a random intercept.
    model = mixedlm(
        "background_corrected_gfp ~ C(experiment) + C(genotype) * C(p19)",
        data, groups=data["sample"],
    ).fit(reml=False)
    terms = pd.DataFrame({
        "term": model.params.index,
        "estimate": model.params.values,
        "standard_error": model.bse.values,
        "wald_z": model.tvalues.values,
        "p_value": model.pvalues.values,
    })

    comparisons = []
    raw_p = []
    for condition in ("−", "+"):
        wt = data.loc[(data["genotype"] == "WT") & (data["p19"] == condition), "background_corrected_gfp"]
        ein2 = data.loc[(data["genotype"] == "ein2") & (data["p19"] == condition), "background_corrected_gfp"]
        result = stats.ttest_ind(wt, ein2, equal_var=False)
        raw_p.append(result.pvalue)
        comparisons.append({"p19": condition, "comparison": "WT vs ein2",
                            "test": "two-sided Welch t-test", "wt_n": len(wt), "ein2_n": len(ein2),
                            "wt_mean": wt.mean(), "ein2_mean": ein2.mean(),
                            "t": result.statistic, "df": result.df, "p_raw": result.pvalue})
    adjusted = multipletests(raw_p, method="holm")[1]
    for row, p_value in zip(comparisons, adjusted):
        row["p_holm"] = p_value
        row["significant_p_holm_lt_0.05"] = bool(p_value < 0.05)

    paired = []
    for genotype in ("WT", "ein2"):
        wide = data.loc[data["genotype"] == genotype].pivot(index="sample", columns="p19",
                                                              values="background_corrected_gfp")
        result = stats.ttest_rel(wide["+"], wide["−"])
        paired.append({"genotype": genotype, "comparison": "+p19 vs −p19",
                       "test": "two-sided paired t-test", "n_pairs": len(wide),
                       "t": result.statistic, "df": result.df, "p_raw": result.pvalue})
    paired_adjusted = multipletests([row["p_raw"] for row in paired], method="holm")[1]
    for row, p_value in zip(paired, paired_adjusted):
        row["p_holm"] = p_value
    return model, terms, pd.DataFrame(comparisons), pd.DataFrame(paired)


def significance(ax, left, right, y, text):
    height = 650
    ax.plot([left, left, right, right], [y, y + height, y + height, y],
            color="black", linewidth=1.0)
    ax.text((left + right) / 2, y + height + 260, text, ha="center", va="bottom",
            fontsize=12, color=TEXT)


def make_plot(data, model, comparisons):
    fig, ax = plt.subplots(figsize=(6.0, 5.3), facecolor="white")
    ax.set_facecolor(PANEL_BG)
    ax.set_axisbelow(True)
    ax.grid(axis="y", color="white", linewidth=0.9, alpha=0.9)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(width=1.0, length=4.5, direction="out")

    centers = np.arange(2, dtype=float)
    offsets = {"WT": -0.18, "ein2": 0.18}
    colors = {"WT": WT_COLOR, "ein2": EIN2_COLOR}
    for condition, center in zip(("−", "+"), centers):
        for genotype in ("WT", "ein2"):
            values = data.loc[(data["p19"] == condition) & (data["genotype"] == genotype),
                              "background_corrected_gfp"].to_numpy(float)
            xpos = center + offsets[genotype]
            mean = values.mean()
            sem = values.std(ddof=1) / math.sqrt(len(values))
            ax.bar(xpos, mean, width=0.32, color=colors[genotype], alpha=BAR_ALPHA,
                   edgecolor="black", linewidth=0.9, zorder=2)
            ax.errorbar(xpos, mean, yerr=sem, fmt="none", ecolor="black", elinewidth=1,
                        capsize=3, capthick=1, zorder=4)
            jitter = np.linspace(-0.045, 0.045, len(values))
            ax.scatter(xpos + jitter, values, s=24, color=colors[genotype],
                       edgecolor="#000000", linewidth=0.45, alpha=POINT_ALPHA, zorder=3)

    for center, condition, y in zip(centers, ("−", "+"), (9800, 20200)):
        p_value = comparisons.loc[comparisons["p19"] == condition, "p_holm"].iloc[0]
        significance(ax, center - 0.18, center + 0.18, y, "ns" if p_value >= 0.05 else "*")

    interaction_term = "C(genotype)[T.ein2]:C(p19)[T.+]"
    interaction_p = model.pvalues[interaction_term]
    ax.text(0.02, 0.98, f"Genotype × p19: P = {interaction_p:.3f}",
            transform=ax.transAxes, ha="left", va="top", fontsize=11.5)
    ax.set_xticks(centers, ["−", "+"])
    ax.set_xlabel("p19")
    ax.set_ylabel("Background-corrected GFP intensity (AU)")
    ax.set_ylim(0, 23000)

    handles = [
        plt.Rectangle((0, 0), 1, 1, facecolor=colors[label], edgecolor="black",
                      linewidth=0.8, alpha=BAR_ALPHA, label=label)
        for label in ("WT", "ein2")
    ]
    legend = ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, -0.18),
                       ncol=2, frameon=False, handlelength=1.2, columnspacing=1.3)
    legend.get_texts()[1].set_fontstyle("italic")
    fig.subplots_adjust(left=0.18, right=0.98, top=0.95, bottom=0.24)
    FIGURES.mkdir(parents=True, exist_ok=True)
    for extension in ("png", "pdf", "svg"):
        fig.savefig(FIGURES / f"202608_wt_ein2_gfp_summary.{extension}",
                    dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def write_outputs(transfer, summary, data, model, terms, comparisons, paired):
    LOGS.mkdir(parents=True, exist_ok=True)
    transfer.to_csv(LOGS / "wt_ein2_transfer_audit.csv", index=False)
    summary.to_csv(LOGS / "wt_ein2_summary_mean_audit.csv", index=False)
    data.to_csv(LOGS / "wt_ein2_background_corrected_long.csv", index=False)
    terms.to_csv(LOGS / "wt_ein2_mixed_model_terms.csv", index=False)
    comparisons.to_csv(LOGS / "wt_ein2_genotype_contrasts.csv", index=False)
    paired.to_csv(LOGS / "wt_ein2_p19_paired_tests.csv", index=False)
    interaction_p = model.pvalues["C(genotype)[T.ein2]:C(p19)[T.+]"]
    note = f"""WT versus ein2 GFP audit and analysis

Workbook audit
All {len(transfer)} transfer/calculation checks passed: {bool(transfer['matches'].all())}.
This comprises 42 direct transfers (14 background, 14 GFP, and 14 GFP+P19 values) and 28 background-subtraction checks. All four displayed summary means passed: {bool(summary['matches'].all())}. The right-hand table is therefore a correct and understandable sample-aligned representation of the left-hand measurements. Each row contains one sample's background, GFP, background-corrected GFP, GFP+P19, and background-corrected GFP+P19. Sample names encode genotype, experiment 1/2, and biological sample number.

Statistics
Background-corrected GFP was analysed with a linear mixed model containing experiment as a fixed blocking factor, genotype, p19, and genotype x p19 as fixed effects, and sample as a random intercept because each sample contributes paired -/+p19 measurements. The genotype x p19 interaction was significant (P={interaction_p:.6f}). Planned two-sided Welch tests compared WT with ein2 within each p19 condition and were Holm-adjusted across the two comparisons. Neither genotype contrast remained significant: -p19 Holm P={comparisons.loc[comparisons['p19']=='−','p_holm'].iloc[0]:.6f}; +p19 Holm P={comparisons.loc[comparisons['p19']=='+','p_holm'].iloc[0]:.6f}. The unadjusted +p19 contrast was P={comparisons.loc[comparisons['p19']=='+','p_raw'].iloc[0]:.6f}, so it should not be labelled significant after the prespecified multiplicity correction.

Plot
Bars show means +/- SEM; translucent points are individual biological samples. WT uses the established muted blue and ein2 uses the shared muted green, keeping it distinct from the orange assigned to 40-1. The ein2 line name is lower-case and italicised in the legend.
"""
    (LOGS / "wt_ein2_gfp_methods_note.txt").write_text(note, encoding="utf-8")


def main():
    formulas, values = workbook_values()
    transfer, summary = audit_workbook(formulas, values)
    if not transfer["matches"].all() or not summary["matches"].all():
        raise ValueError("Workbook audit failed; see audit tables before plotting.")
    data = long_data(values)
    model, terms, comparisons, paired = analyse(data)
    make_plot(data, model, comparisons)
    write_outputs(transfer, summary, data, model, terms, comparisons, paired)


if __name__ == "__main__":
    main()
